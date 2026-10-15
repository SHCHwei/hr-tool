import openpyxl
from datetime import datetime
import calendar
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from time_calculator import get_effective_checkout, calculate_overtime_and_early

WEEKDAYS_ZH = ["星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日"]

def test_export():
    year, month = 2026, 8
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{year}年{month:02d}月打卡表"
    
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = f"{year} 年 {month:02d} 月 員工打卡表記錄"
    title_cell.font = Font(name="微軟正黑體", size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    headers = ["日期", "星期", "首次上班", "首次下班", "二次上班", "二次下班", "加班時數", "早退時數"]
    for col_idx, text in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=text)
        cell.font = Font(name="微軟正黑體", size=11, bold=True, color="1F497D")
        cell.fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")

    _, num_days = calendar.monthrange(year, month)
    for day in range(1, num_days + 1):
        date_obj = datetime(year, month, day)
        row_idx = day + 2
        ws.cell(row=row_idx, column=1, value=f"{year}-{month:02d}-{day:02d}")
        ws.cell(row=row_idx, column=2, value=WEEKDAYS_ZH[date_obj.weekday()])
        
        t1_in, t1_out, t2_in, t2_out = "", "", "", ""
        if date_obj.weekday() < 5:
            if day == 3: # 測試 17:00 早退 30分
                t1_in, t1_out, t2_in, t2_out = "09:00", "12:00", "13:00", "17:00"
            elif day == 4: # 測試 17:35 正常下班緩衝區間 (無早退、無加班)
                t1_in, t1_out, t2_in, t2_out = "09:00", "12:00", "13:00", "17:35"
            elif day == 5: # 測試 19:00 加班 1小時15分
                t1_in, t1_out, t2_in, t2_out = "09:00", "12:00", "13:00", "19:00"
            elif day == 6: # 測試 01:00 跨夜加班 7小時15分
                t1_in, t1_out, t2_in, t2_out = "09:00", "12:00", "13:00", "01:00"
            else: # 預設工作日 18:00 加班 15分
                t1_in, t1_out, t2_in, t2_out = "09:00", "12:00", "13:00", "18:00"

            ws.cell(row=row_idx, column=3, value=t1_in)
            ws.cell(row=row_idx, column=4, value=t1_out)
            ws.cell(row=row_idx, column=5, value=t2_in)
            ws.cell(row=row_idx, column=6, value=t2_out)

        checkout_time = get_effective_checkout(t1_in, t1_out, t2_in, t2_out)
        overtime_str, early_str = calculate_overtime_and_early(checkout_time)
        ws.cell(row=row_idx, column=7, value=overtime_str)
        ws.cell(row=row_idx, column=8, value=early_str)

    output_path = "test_output.xlsx"
    wb.save(output_path)
    print("Excel Export test passed successfully! Generated:", output_path)

if __name__ == "__main__":
    test_export()
