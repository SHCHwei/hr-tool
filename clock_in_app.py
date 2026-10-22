import sys
import calendar
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from time_calculator import calculate_daily_overtime_and_early, calculate_late, is_valid_hhmm

from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QComboBox, QLineEdit, QPushButton, QScrollArea, QFrame,
    QGridLayout, QFileDialog, QMessageBox, QGroupBox, QSpinBox
)
from PySide6.QtCore import Qt, QSize, QRegularExpression
from PySide6.QtGui import QFont, QColor, QPalette, QRegularExpressionValidator

# 星期中文對照表
WEEKDAYS_ZH = ["星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日"]


def create_light_palette():
    palette = QPalette()
    color_roles = {
        QPalette.Window: "#F4F7F6",
        QPalette.WindowText: "#333333",
        QPalette.Base: "#FFFFFF",
        QPalette.AlternateBase: "#F4F7F6",
        QPalette.ToolTipBase: "#FFFFFF",
        QPalette.ToolTipText: "#333333",
        QPalette.Text: "#222222",
        QPalette.Button: "#FFFFFF",
        QPalette.ButtonText: "#333333",
        QPalette.BrightText: "#C62828",
        QPalette.Link: "#357ABD",
        QPalette.LinkVisited: "#5B6F9C",
        QPalette.Highlight: "#4A90E2",
        QPalette.HighlightedText: "#FFFFFF",
        QPalette.PlaceholderText: "#888888",
    }
    disabled_roles = {
        QPalette.WindowText,
        QPalette.Text,
        QPalette.ButtonText,
        QPalette.HighlightedText,
    }

    for role, color in color_roles.items():
        palette.setColor(QPalette.Active, role, QColor(color))
        palette.setColor(QPalette.Inactive, role, QColor(color))
        disabled_color = "#888888" if role in disabled_roles else color
        palette.setColor(QPalette.Disabled, role, QColor(disabled_color))

    return palette


class ClockInApp(QMainWindow):
    def __init__(self, employee_name="", start_time="", end_time="", on_close=None):
        super().__init__()
        self.employee_name = employee_name
        self.start_time = start_time
        self.end_time = end_time
        self.on_close = on_close
        self.setWindowTitle("月打卡記錄表輸入系統 (Clock-in Management)")
        self.resize(850, 750)
        self.setMinimumSize(750, 600)
        
        # 儲存每日輸入框的字典: {day_num: [edit_1in, edit_1out, edit_2in, edit_2out]}
        self.daily_inputs = {}
        self.daily_results = {}
        
        self.init_ui()
        self.apply_styles()
        self.on_date_changed()

    def closeEvent(self, event):
        super().closeEvent(event)
        if event.isAccepted() and self.on_close:
            self.on_close()

    def init_ui(self):
        main_widget = QWidget()
        self.setCentralWidget(main_widget)
        main_layout = QVBoxLayout(main_widget)
        main_layout.setSpacing(12)
        main_layout.setContentsMargins(15, 15, 15, 15)

        # -------------------------------------------------------------
        # 1. 頂部控制面板 (年份、月份選擇)
        # -------------------------------------------------------------
        top_group = QGroupBox("月曆選擇")
        top_layout = QVBoxLayout(top_group)
        controls_layout = QHBoxLayout()
        controls_layout.setSpacing(15)

        # 年份選擇
        controls_layout.addWidget(QLabel("年份:"))
        self.year_spin = QSpinBox()
        current_year = datetime.now().year
        self.year_spin.setRange(2000, 2040)
        self.year_spin.setValue(current_year)
        self.year_spin.valueChanged.connect(self.on_date_changed)
        controls_layout.addWidget(self.year_spin)

        # 月份選擇
        controls_layout.addWidget(QLabel("月份:"))
        self.month_spin = QSpinBox()
        current_month = datetime.now().month
        self.month_spin.setRange(1, 12)
        self.month_spin.setValue(current_month)
        self.month_spin.valueChanged.connect(self.on_date_changed)
        controls_layout.addWidget(self.month_spin)

        controls_layout.addStretch(1)


        # 返回
        self.btn_back = QPushButton("返回")
        self.btn_back.clicked.connect(self.close)
        controls_layout.addWidget(self.btn_back)


        # 清空按鈕
        self.btn_clear = QPushButton("🗑️ 清空表格")
        self.btn_clear.clicked.connect(self.clear_all_inputs)
        controls_layout.addWidget(self.btn_clear)

        employee_layout = QHBoxLayout()
        employee_layout.addWidget(QLabel("員工姓名:"))
        self.employee_name_label = QLabel(self.employee_name)
        employee_layout.addWidget(self.employee_name_label, stretch=1)
        employee_layout.addStretch(1)
        self.btn_import = QPushButton("匯入")
        employee_layout.addWidget(self.btn_import)

        top_layout.addLayout(controls_layout)
        top_layout.addLayout(employee_layout)

        main_layout.addWidget(top_group)

        # -------------------------------------------------------------
        # 2. 中間打卡表格標題列
        # -------------------------------------------------------------
        header_frame = QFrame()
        header_frame.setObjectName("headerFrame")
        header_layout = QGridLayout(header_frame)
        header_layout.setContentsMargins(10, 8, 25, 8) # 25px right margin to align with scrollbar

        lbl_date = QLabel("日期 / 星期")
        lbl_date.setAlignment(Qt.AlignCenter)
        lbl_date.setStyleSheet("font-weight: bold; color: #2C3E50;")
        header_layout.addWidget(lbl_date, 0, 0)

        headers = ["首次上班", "首次下班", "二次上班", "二次下班", "加班時數", "早退時數", "遲到時數"]
        for col_idx, text in enumerate(headers, start=1):
            lbl = QLabel(text)
            lbl.setAlignment(Qt.AlignCenter)
            lbl.setStyleSheet("font-weight: bold; color: #2C3E50;")
            header_layout.addWidget(lbl, 0, col_idx)

        header_layout.setColumnStretch(0, 2)
        header_layout.setColumnStretch(1, 2)
        header_layout.setColumnStretch(2, 2)
        header_layout.setColumnStretch(3, 2)
        header_layout.setColumnStretch(4, 2)
        header_layout.setColumnStretch(5, 2)
        header_layout.setColumnStretch(6, 2)
        header_layout.setColumnStretch(7, 2)

        main_layout.addWidget(header_frame)

        # -------------------------------------------------------------
        # 3. 中間滾動區域 (顯示1號到最後一號的每日輸入框)
        # -------------------------------------------------------------
        self.scroll_area = QScrollArea()
        self.scroll_area.setObjectName("tableScrollArea")
        self.scroll_area.setWidgetResizable(True)
        self.scroll_area.setFrameShape(QFrame.NoFrame)
        self.scroll_area.viewport().setStyleSheet("background-color: #FFFFFF;")

        self.scroll_content = QWidget()
        self.scroll_content.setObjectName("scrollContent")
        self.scroll_content.setStyleSheet("background-color: #FFFFFF;")
        self.scroll_layout = QVBoxLayout(self.scroll_content)
        self.scroll_layout.setSpacing(6)
        self.scroll_layout.setContentsMargins(6, 6, 6, 6)
        self.scroll_layout.addStretch(1)

        self.scroll_area.setWidget(self.scroll_content)
        main_layout.addWidget(self.scroll_area, stretch=1)

        # -------------------------------------------------------------
        # 4. 視窗底部：匯出 Excel 按鈕
        # -------------------------------------------------------------
        bottom_frame = QFrame()
        bottom_layout = QHBoxLayout(bottom_frame)
        bottom_layout.setContentsMargins(0, 10, 0, 5)

        self.btn_export = QPushButton("📊 匯出 Excel 打卡表")
        self.btn_export.setFixedHeight(48)
        self.btn_export.setFont(QFont("Microsoft JhengHei", 12, QFont.Bold))
        self.btn_export.clicked.connect(self.export_to_excel)
        bottom_layout.addWidget(self.btn_export)

        main_layout.addWidget(bottom_frame)

    def apply_styles(self):
        """套用現代美觀 CSS 樣式，確保所有輸入框白底黑字、表格全區白底無黑塊"""
        qss = """
        QMainWindow {
            background-color: #F4F7F6;
        }
        QGroupBox {
            /* font-weight: bold; */
            /* font-size: 13px; */
            /* border: 1px solid #DCDFE6; */
            /* border-radius: 8px; */
            /* margin-top: 6px; */
            /* padding-top: 12px; */
            /* background-color: #FFFFFF; */
            /* color: #333333; */
        }
        QGroupBox::title {
            subcontrol-origin: margin;
            left: 12px;
            padding: 0 5px;
            color: #333333;
        }
        QLabel {
            color: #333333;
        }
        /* 年份、月份選擇框：白底黑字 */
        QSpinBox:focus {
            border: 1px solid #4A90E2;
        }
        QSpinBox::up-button, QSpinBox::down-button {
            width: 18px;
            /* background-color: #000000; */
            border-left: 1px solid #555555;
        }
        QSpinBox::up-arrow {
            width: 0;
            height: 0;
            border-left: 4px solid transparent;
            border-right: 4px solid transparent;
            border-bottom: 5px solid #FFFFFF;
        }
        QSpinBox::down-arrow {
            width: 0;
            height: 0;
            border-left: 4px solid transparent;
            border-right: 4px solid transparent;
            border-top: 5px solid #FFFFFF;
        }
        QSpinBox QLineEdit {
            background-color: #FFFFFF;
            color: #000000;
        }
        #headerFrame {
            background-color: #E8EEF5;
            border-radius: 6px;
            border: 1px solid #D0D7DE;
        }
        #headerFrame QLabel {
            color: #2C3E50;
            font-weight: bold;
        }
        /* 表格滾動區域與內容背景：一律為純白底色 */
        #tableScrollArea {
            background-color: #FFFFFF;
            border: 1px solid #D0D7DE;
            border-radius: 6px;
        }
        #tableScrollArea QWidget {
            background-color: #FFFFFF;
        }
        #scrollContent {
            background-color: #FFFFFF;
        }
        /* 所有時間輸入框：白底黑字 */
        QLineEdit {
            border: 1px solid #DCDFE6;
            border-radius: 4px;
            padding: 5px 8px;
            font-size: 13px;
            background-color: #FFFFFF;
            color: #000000;
            selection-background-color: #4A90E2;
            selection-color: #FFFFFF;
        }
        QLineEdit:focus {
            border: 1px solid #4A90E2;
            background-color: #F0F7FF;
            color: #000000;
        }
        QPushButton {
            background-color: #4A90E2;
            color: white;
            border-radius: 6px;
            padding: 6px 14px;
            font-size: 13px;
            font-weight: bold;
            border: none;
        }
        QPushButton:hover {
            background-color: #357ABD;
        }
        QPushButton:pressed {
            background-color: #2D68A4;
        }
        #btn_export {
            background-color: #27AE60;
        }
        #btn_export:hover {
            background-color: #219653;
        }
        #btn_export:pressed {
            background-color: #1E8449;
        }
        """
        self.setStyleSheet(qss)
        self.btn_export.setObjectName("btn_export")

    def on_date_changed(self):
        """當年月改變時重新生成打卡表格內容"""
        year = self.year_spin.value()
        month = self.month_spin.value()

        # 清除舊項目
        for i in reversed(range(self.scroll_layout.count())):
            widget = self.scroll_layout.itemAt(i).widget()
            if widget:
                widget.setParent(None)

        self.daily_inputs.clear()
        self.daily_results.clear()
        
        # 計算該月份的天數 (例如 28, 29, 30, 31 天)
        _, num_days = calendar.monthrange(year, month)

        # 逐日建立輸入列
        for day in range(1, num_days + 1):
            date_obj = datetime(year, month, day)
            weekday_idx = date_obj.weekday() # 0: 週一, ..., 6: 週日
            weekday_str = WEEKDAYS_ZH[weekday_idx]
            
            row_frame = QFrame()
            row_frame.setObjectName("rowFrame")
            row_frame.setFrameShape(QFrame.StyledPanel)
            
            # 設定週末與工作日的背景顏色區隔
            if weekday_idx == 5: # 週六
                row_frame.setStyleSheet("#rowFrame { background-color: #F0F7FF; border: 1px solid #D0E3F8; border-radius: 6px; }")
            elif weekday_idx == 6: # 週日
                row_frame.setStyleSheet("#rowFrame { background-color: #FFF5F5; border: 1px solid #FCD4D4; border-radius: 6px; }")
            else:
                row_frame.setStyleSheet("#rowFrame { background-color: #FFFFFF; border: 1px solid #E4E7ED; border-radius: 6px; }")

            grid = QGridLayout(row_frame)
            grid.setContentsMargins(10, 6, 10, 6)
            grid.setSpacing(8)

            # 日期標籤 
            date_display = f"{month:02d}月{day:02d}日 ({weekday_str})"
            lbl_day = QLabel(date_display)
            lbl_day.setFont(QFont("Microsoft JhengHei", 10))
            if weekday_idx >= 5:
                lbl_day.setStyleSheet("background: transparent; font-weight: bold; color: #E67E22;" if weekday_idx == 5 else "background: transparent; font-weight: bold; color: #E74C3C;")
            else:
                lbl_day.setStyleSheet("background: transparent; color: #333333;")
            
            grid.addWidget(lbl_day, 0, 0)

            # 四格打卡輸入框
            inputs = []
            placeholders = [self.start_time or "0830", "1200", "1300", self.end_time or "1730"]
            for col in range(4):
                edit = QLineEdit()
                edit.setPlaceholderText(placeholders[col])
                edit.setAlignment(Qt.AlignCenter)
                edit.setValidator(QRegularExpressionValidator(QRegularExpression("[0-9]{0,4}"), edit))
                grid.addWidget(edit, 0, col + 1)
                edit.textChanged.connect(lambda _text, day=day: self.update_daily_calculation(day))
                edit.editingFinished.connect(lambda edit=edit: self.validate_time_input(edit))
                inputs.append(edit)

            overtime_label = QLabel()
            early_label = QLabel()
            late_label = QLabel()
            overtime_label.setAlignment(Qt.AlignCenter)
            early_label.setAlignment(Qt.AlignCenter)
            late_label.setAlignment(Qt.AlignCenter)
            grid.addWidget(overtime_label, 0, 5)
            grid.addWidget(early_label, 0, 6)
            grid.addWidget(late_label, 0, 7)

            # 寬度比例分配
            grid.setColumnStretch(0, 2)
            grid.setColumnStretch(1, 2)
            grid.setColumnStretch(2, 2)
            grid.setColumnStretch(3, 2)
            grid.setColumnStretch(4, 2)
            grid.setColumnStretch(5, 2)
            grid.setColumnStretch(6, 2)
            grid.setColumnStretch(7, 2)

            self.daily_inputs[day] = inputs
            self.daily_results[day] = (overtime_label, early_label, late_label)
            self.scroll_layout.addWidget(row_frame)

        self.scroll_layout.addStretch(1)

    def validate_time_input(self, edit):
        value = edit.text()
        if value and not is_valid_hhmm(value):
            edit.clear()

    def update_daily_calculation(self, day):
        inputs = self.daily_inputs.get(day)
        results = self.daily_results.get(day)
        if not inputs or not results:
            return

        year = self.year_spin.value()
        month = self.month_spin.value()
        is_weekend = datetime(year, month, day).weekday() >= 5
        overtime_str, early_str = calculate_daily_overtime_and_early(
            *(edit.text() for edit in inputs),
            is_weekend=is_weekend,
            work_off=self.end_time,
            work_start=self.start_time,
        )
        results[0].setText(overtime_str)
        results[1].setText(early_str)
        results[2].setText("" if is_weekend else calculate_late(inputs[0].text(), self.start_time))

    def clear_all_inputs(self):
        """清空所有打卡輸入格"""
        for day, inputs in self.daily_inputs.items():
            for edit in inputs:
                edit.clear()

    def export_to_excel(self):
        """將打卡資料匯出為造型美觀的 Excel (.xlsx) 檔案"""
        year = self.year_spin.value()
        month = self.month_spin.value()
        
        employee_name = self.employee_name_label.text().strip()
        filename_prefix = f"{employee_name}_" if employee_name else ""
        default_filename = f"{filename_prefix}員工打卡表記錄_{year}年{month:02d}月.xlsx"
        file_path, _ = QFileDialog.getSaveFileName(
            self, "儲存打卡表 Excel", default_filename, "Excel Files (*.xlsx)"
        )
        
        if not file_path:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = f"{year}年{month:02d}月打卡表"
            ws.views.sheetView[0].showGridLines = True

            # ---------------------------------------------------------
            # 1. Excel 大標題
            # ---------------------------------------------------------
            ws.merge_cells("A1:I1")
            title_cell = ws["A1"]
            title_cell.value = f"{year} 年 {month:02d} 月 員工打卡表記錄"
            title_cell.font = Font(name="微軟正黑體", size=16, bold=True, color="FFFFFF")
            title_cell.fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            title_cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 40

            # ---------------------------------------------------------
            # 2. 表格欄位標題 (Header)
            # ---------------------------------------------------------
            headers = ["日期", "星期", "首次上班", "首次下班", "二次上班", "二次下班", "加班時數", "早退時數", "遲到時數"]
            ws.row_dimensions[2].height = 25
            
            header_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
            header_font = Font(name="微軟正黑體", size=11, bold=True, color="1F497D")
            thin_border = Border(
                left=Side(style='thin', color='BFBFBF'),
                right=Side(style='thin', color='BFBFBF'),
                top=Side(style='thin', color='BFBFBF'),
                bottom=Side(style='thin', color='BFBFBF')
            )

            for col_idx, header_text in enumerate(headers, start=1):
                cell = ws.cell(row=2, column=col_idx, value=header_text)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

            # ---------------------------------------------------------
            # 3. 填入每日資料列
            # ---------------------------------------------------------
            weekend_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
            
            for day, inputs in self.daily_inputs.items():
                date_obj = datetime(year, month, day)
                weekday_idx = date_obj.weekday()
                weekday_str = WEEKDAYS_ZH[weekday_idx]
                
                date_str = f"{year}-{month:02d}-{day:02d}"
                t1_in = inputs[0].text().strip()
                t1_out = inputs[1].text().strip()
                t2_in = inputs[2].text().strip()
                t2_out = inputs[3].text().strip()

                overtime_str, early_str = calculate_daily_overtime_and_early(
                    t1_in, t1_out, t2_in, t2_out,
                    is_weekend=weekday_idx >= 5,
                    work_off=self.end_time,
                    work_start=self.start_time,
                )

                late_str = "" if weekday_idx >= 5 else calculate_late(t1_in, self.start_time)

                row_idx = day + 2 # row 1: title, row 2: header
                ws.row_dimensions[row_idx].height = 22
                
                row_data = [date_str, weekday_str, t1_in, t1_out, t2_in, t2_out, overtime_str, early_str, late_str]

                for col_idx, val in enumerate(row_data, start=1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=val)
                    cell.font = Font(name="微軟正黑體", size=10)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = thin_border

                    # 如果是週末，加入淡淡的背景色區隔
                    if weekday_idx >= 5:
                        cell.fill = weekend_fill
                        if col_idx == 2: # 星期標註顏色
                            cell.font = Font(name="微軟正黑體", size=10, bold=True, color="C00000")

            # ---------------------------------------------------------
            # 4. 自動調整欄寬
            # ---------------------------------------------------------
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or '')
                    # 計算長度（中文算雙倍長度）
                    cell_len = sum(2 if ord(c) > 127 else 1 for c in val_str)
                    if cell_len > max_len:
                        max_len = cell_len
                ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

            # 儲存檔案
            wb.save(file_path)
            
            QMessageBox.information(
                self, "匯出成功", f"打卡表已成功匯出至：\n{file_path}"
            )

        except Exception as e:
            QMessageBox.critical(
                self, "匯出失敗", f"匯出 Excel 時發生錯誤：\n{str(e)}"
            )


def main():
    app = QApplication(sys.argv)
    
    # 不依賴作業系統主題，明確套用所有狀態的淺色 Palette
    app.setPalette(create_light_palette())

    # 設置字體
    font = QFont("Microsoft JhengHei", 10)
    app.setFont(font)

    window = ClockInApp()
    window.show()
    sys.exit(app.exec())

if __name__ == "__main__":
    main()
